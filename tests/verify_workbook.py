"""Read-only verification of artifact-tool export against typed payload."""
import json
import math
from pathlib import Path
from openpyxl import load_workbook

root=Path(__file__).resolve().parents[1]
out=root/'outputs'/'pvdelta_v1'
source=json.loads((out/'workbook_payload.json').read_text(encoding='utf-8'))
workbook=load_workbook(out/'重整数据与审计.xlsx',read_only=True,data_only=True)
checked=0
errors=[]
for name,obj in source.items():
    matrix=obj if isinstance(obj,list) else [obj['columns']]+obj['data']
    sheet=workbook[name]
    print(name,sheet.max_row,sheet.max_column,flush=True)
    for i,row in enumerate(sheet.iter_rows(min_row=1,max_row=len(matrix),max_col=len(matrix[0]),values_only=True)):
        for j,actual in enumerate(row):
            expected=matrix[i][j]
            if isinstance(expected,(int,float)) and not isinstance(expected,bool) and isinstance(actual,(int,float)):
                good=math.isclose(float(actual),float(expected),rel_tol=1e-9,abs_tol=1e-8)
            else:
                good=actual==expected
            if not good:
                errors.append((name,i+1,j+1,str(expected),str(actual)))
            checked+=1
print('Verified cells:',checked,'mismatches:',len(errors),flush=True)
print(errors[:5],flush=True)
assert not errors
workbook.close()
